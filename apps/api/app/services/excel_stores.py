from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.db.models import Store
from app.domain.store_flow import TZ_TASHKENT, fmt_store_date

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill(fill_type="solid", start_color="305496")
_BODY_FONT = Font(color="000000")
_BODY_FILL = PatternFill(fill_type="solid", start_color="FFFFFF")
_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _style_sheet(ws: Worksheet) -> None:
    """Rang va fillarni har bir katakka aniq yozamiz.

    Fill berilmagan katakni ba'zi telefon ko'ruvchilari (to'q rejimda) qora fon
    bilan chizadi, yozuv esa qora — jadval umuman ko'rinmay qoladi.
    """
    sum_cols: set[int] = set()
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.border = _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if isinstance(cell.value, str) and "so'm" in cell.value:
            sum_cols.add(cell.column)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = _BODY_FONT
            cell.fill = _BODY_FILL
            cell.border = _BORDER
            if cell.column in sum_cols and isinstance(cell.value, int):
                cell.number_format = "#,##0"

    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        longest = max(
            (len(str(c.value)) for c in ws[letter] if c.value is not None), default=0
        )
        ws.column_dimensions[letter].width = min(max(longest + 2, 10), 40)

    ws.freeze_panes = "A2"


def stores_to_xlsx_bytes(
    stores: list[Store],
    global_tok_price_per_kw: int | None = None,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Magazinlar"
    ws.append(
        [
            "ID",
            "Nomi",
            "Telefon",
            "Manzil",
            "Hisobot sanasi",
            "Oylik (so'm)",
            "Qarz (so'm)",
            "Tok kW (joriy)",
            "Tok qarz / oxirgi +kW",
            "Umumiy tok narxi (so'm/kW)",
            "Tok to'lov (so'm)",
            "Jami qarz (so'm)",
            "Yaratilgan",
        ]
    )
    for s in stores:
        created = "—"
        if s.created_at:
            created = s.created_at.astimezone(TZ_TASHKENT).strftime("%d.%m.%Y %H:%M")
        tok_pay = ""
        if global_tok_price_per_kw is not None:
            tok_pay = int(s.debt_tok or 0) * int(global_tok_price_per_kw)
        total_debt = int(s.debt_balance or 0) + (tok_pay if isinstance(tok_pay, int) else 0)
        ws.append(
            [
                s.id,
                s.name,
                s.owner_phone or "",
                s.address or "",
                fmt_store_date(s.store_date),
                s.monthly_amount if s.monthly_amount is not None else "",
                int(s.debt_balance or 0),
                s.electricity_kw if s.electricity_kw is not None else "",
                int(s.debt_tok or 0),
                int(global_tok_price_per_kw) if global_tok_price_per_kw is not None else "",
                tok_pay,
                total_debt,
                created,
            ]
        )

    _style_sheet(ws)
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _fmt_ts(ca: datetime) -> str:
    if ca.tzinfo is None:
        ca = ca.replace(tzinfo=TZ_TASHKENT)
    return ca.astimezone(TZ_TASHKENT).strftime("%d.%m.%Y %H:%M")


def admin_report_xlsx_bytes(
    stores: list[Store],
    payment_rows: list[
        tuple[int, int, str, int, int, datetime, int]
    ],  # id, sid, name, amount, debt_after, created_at, admin_id
    electricity_rows: list[
        tuple[int, int, str, datetime, datetime, int, int, int, datetime]
    ],  # log_id, sid, name, p_from, p_to, rb, ra, delta, created_at
    global_tok_price_per_kw: int | None = None,
) -> bytes:
    """Bitta fayl: magazinlar + qarzdan ayirishlar + tok tarixi (kW)."""
    wb = Workbook()
    ws1 = wb.active
    assert ws1 is not None
    ws1.title = "Magazinlar"
    ws1.append(
        [
            "ID",
            "Nomi",
            "Telefon",
            "Manzil",
            "Hisobot sanasi",
            "Oylik (so'm)",
            "Qarz (so'm)",
            "Tok kW (joriy)",
            "Tok qarz / oxirgi +kW",
            "Umumiy tok narxi (so'm/kW)",
            "Tok to'lov (so'm)",
            "Jami qarz (so'm)",
            "Yaratilgan",
        ]
    )
    for s in stores:
        created = "—"
        if s.created_at:
            created = s.created_at.astimezone(TZ_TASHKENT).strftime("%d.%m.%Y %H:%M")
        tok_pay = ""
        if global_tok_price_per_kw is not None:
            tok_pay = int(s.debt_tok or 0) * int(global_tok_price_per_kw)
        total_debt = int(s.debt_balance or 0) + (tok_pay if isinstance(tok_pay, int) else 0)
        ws1.append(
            [
                s.id,
                s.name,
                s.owner_phone or "",
                s.address or "",
                fmt_store_date(s.store_date),
                s.monthly_amount if s.monthly_amount is not None else "",
                int(s.debt_balance or 0),
                s.electricity_kw if s.electricity_kw is not None else "",
                int(s.debt_tok or 0),
                int(global_tok_price_per_kw) if global_tok_price_per_kw is not None else "",
                tok_pay,
                total_debt,
                created,
            ]
        )

    ws2 = wb.create_sheet("Qarzdan ayirishlar")
    ws2.append(
        [
            "ID",
            "Magazin ID",
            "Magazin nomi",
            "Sana vaqt",
            "Ayirilgan summa (so'm)",
            "Qarzdan keyin (so'm)",
            "Admin Telegram ID",
        ]
    )
    for row in payment_rows:
        pid, sid, name, amount, debt_after, created_at, admin_id = row
        ts = _fmt_ts(created_at)
        ws2.append([pid, sid, name, ts, amount, debt_after, admin_id])

    ws3 = wb.create_sheet("Tok tarixi (kW)")
    ws3.append(
        [
            "ID",
            "Magazin ID",
            "Magazin nomi",
            "Davr boshlanishi",
            "Davr tugashi",
            "Eski ko'rsatkich (kW)",
            "Yangi ko'rsatkich (kW)",
            "Iste'mol (+kW)",
            "Yozilgan vaqt",
        ]
    )
    for row in electricity_rows:
        lid, sid, name, p_from, p_to, rb, ra, delta, created_at = row
        ws3.append(
            [
                lid,
                sid,
                name,
                _fmt_ts(p_from),
                _fmt_ts(p_to),
                rb,
                ra,
                delta,
                _fmt_ts(created_at),
            ]
        )

    for sheet in (ws1, ws2, ws3):
        _style_sheet(sheet)
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def single_store_electricity_excel_bytes(
    store_id: int,
    store_name: str,
    rows: list[
        tuple[int, datetime, datetime, int, int, int, datetime]
    ],  # id, p_from, p_to, rb, ra, delta, created_at
) -> bytes:
    """Bitta magazin — tok (kW) o'zgarishlari."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Tok_tarixi"
    ws.append(
        [
            "Magazin ID",
            "Magazin",
            "Yozuv ID",
            "Davr boshlanishi",
            "Davr tugashi",
            "Eski ko'rsatkich (kW)",
            "Yangi ko'rsatkich (kW)",
            "Iste'mol (+kW)",
            "Yozilgan vaqt",
        ]
    )
    for lid, p_from, p_to, rb, ra, delta, created_at in rows:
        ws.append(
            [
                store_id,
                store_name,
                lid,
                _fmt_ts(p_from),
                _fmt_ts(p_to),
                rb,
                ra,
                delta,
                _fmt_ts(created_at),
            ]
        )

    _style_sheet(ws)
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def single_store_debt_payments_excel_bytes(
    store_id: int,
    store_name: str,
    rows: list[tuple[int, int, int, datetime, int]],  # id, amount, debt_after, created_at, admin_id
) -> bytes:
    """Bitta magazin — qarzdan ayirishlar (to'lovlar)."""
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Tolovlar"
    ws.append(
        [
            "Magazin ID",
            "Magazin",
            "Yozuv ID",
            "Sana vaqt",
            "Ayirilgan summa (so'm)",
            "Qarzdan keyin (so'm)",
            "Admin Telegram ID",
        ]
    )
    for pid, amount, debt_after, created_at, admin_id in rows:
        ws.append(
            [
                store_id,
                store_name,
                pid,
                _fmt_ts(created_at),
                amount,
                debt_after,
                admin_id,
            ]
        )

    _style_sheet(ws)
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()
